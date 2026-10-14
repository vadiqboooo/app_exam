"""Readers for the school import wizard: teachers, groups and the previous app version.

Like ``excel.py`` this module is an adapter only: it never touches the new database.
"""

import os
import re
import sqlite3
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.importers.group_details import group_details
from app.importers.record import StudentRecord, normalize_name

SQLITE_MAGIC = b"SQLite format 3\x00"


@dataclass(frozen=True)
class TeacherRecord:
    name: str
    subjects: tuple[str, ...] = ()


@dataclass(frozen=True)
class GroupRecord:
    name: str
    subject: str | None
    exam_format: str | None
    teacher: str | None


@dataclass(frozen=True)
class RosterData:
    """Normalized contents of one or more group-roster exports."""

    teachers: tuple[TeacherRecord, ...]
    groups: tuple[GroupRecord, ...]
    students: tuple[StudentRecord, ...]


@dataclass(frozen=True)
class LegacyWork:
    fio: str
    subject_code: str
    scores: tuple[float | None, ...]
    comment: str


@dataclass(frozen=True)
class LegacyData:
    title: str | None
    created_at: datetime | None
    works: tuple[LegacyWork, ...]


def _table(source: BinaryIO) -> tuple[list[str], list[tuple]]:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise ValueError("Не удалось открыть XLSX") from error
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [normalize_name(str(value or "")) for value in next(rows, ())]
        body = [row for row in rows if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()
    return headers, body


def _find(headers: list[str], *names: str) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def _cell(row: tuple, index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return " ".join(str(row[index]).split())


def read_teachers(source: BinaryIO) -> list[TeacherRecord]:
    headers, rows = _table(source)
    name_col = _find(headers, "фио", "сотрудник", "учитель")
    if name_col is None:
        raise ValueError("В файле учителей нет колонки «ФИО»")
    subject_col = _find(headers, "предметы", "предмет")
    records, seen = [], set()
    for row in rows:
        name = _cell(row, name_col)
        if not name:
            continue
        key = normalize_name(name)
        if key in seen:
            raise ValueError(f"Повтор сотрудника в файле: {name}")
        seen.add(key)
        subjects = tuple(
            dict.fromkeys(
                part.strip()
                for part in re.split(r"[,;\r\n]+", _cell(row, subject_col))
                if part.strip()
            )
        )
        records.append(TeacherRecord(name, subjects))
    if not records:
        raise ValueError("В файле нет сотрудников")
    return records


def read_groups(source: BinaryIO) -> list[GroupRecord]:
    headers, rows = _table(source)
    name_col = _find(headers, "название группы", "группа", "название")
    if name_col is None:
        raise ValueError("В файле групп нет колонки «Название группы»")
    subject_col = _find(headers, "предмет")
    type_col = _find(headers, "тип", "тип группы", "формат")
    teacher_col = _find(headers, "учитель", "преподаватель")
    records, seen = [], set()
    for row in rows:
        name = _cell(row, name_col)
        if not name:
            continue
        if name in seen:
            raise ValueError(f"Повтор группы в файле: {name}")
        seen.add(name)
        details = group_details(name)
        kind = normalize_name(_cell(row, type_col))
        exam_format = "oge" if kind == "огэ" else "ege" if kind == "егэ" else details["exam_format"]
        records.append(
            GroupRecord(
                name=name,
                subject=_cell(row, subject_col) or details["subject"],
                exam_format=exam_format,
                teacher=_cell(row, teacher_col) or details["teacher_name"],
            )
        )
    if not records:
        raise ValueError("В файле нет групп")
    return records


def _grade_from_group(name: str, exam_format: str | None) -> int | None:
    explicit = re.search(r"(?<!\d)(1[01]|[1-9])\s*(?:класс|кл\.?)(?!\w)", name, re.IGNORECASE)
    if explicit:
        return int(explicit.group(1))
    return {"oge": 9, "ege": 11}.get(exam_format or "")


def _pupils(value: str) -> tuple[str, ...]:
    return tuple(
        dict.fromkeys(
            " ".join(part.split()) for part in re.split(r"[,;\r\n]+", value) if part.strip()
        )
    )


def read_group_roster(source: BinaryIO) -> RosterData:
    """Read an export where every row contains a group, its teacher and its pupils."""
    headers, rows = _table(source)
    name_col = _find(headers, "название", "название группы", "группа")
    pupils_col = _find(headers, "клиенты", "ученики", "учащиеся")
    teacher_col = _find(
        headers,
        "отв. педагог",
        "ответственный педагог",
        "ответственный учитель",
        "учитель",
        "преподаватель",
    )
    missing = []
    if name_col is None:
        missing.append("Название")
    if pupils_col is None:
        missing.append("Клиенты")
    if teacher_col is None:
        missing.append("Отв. педагог")
    if missing:
        raise ValueError("В файле нет колонок: " + ", ".join(f"«{name}»" for name in missing))

    groups: list[GroupRecord] = []
    teachers: dict[str, dict] = {}
    students: dict[str, dict] = {}
    seen_groups: set[str] = set()
    for number, row in enumerate(rows, start=2):
        name = _cell(row, name_col)
        if not name:
            continue
        if name in seen_groups:
            raise ValueError(f"Строка {number}: повтор группы в файле: {name}")
        seen_groups.add(name)
        teacher = _cell(row, teacher_col) or None
        details = group_details(name)
        subject = details["subject"]
        exam_format = details["exam_format"]
        groups.append(GroupRecord(name, subject, exam_format, teacher))

        if teacher:
            key = normalize_name(teacher)
            item = teachers.setdefault(key, {"name": teacher, "subjects": []})
            if subject and subject not in item["subjects"]:
                item["subjects"].append(subject)

        grade = _grade_from_group(name, exam_format)
        for pupil in _pupils(_cell(row, pupils_col)):
            key = normalize_name(pupil)
            item = students.setdefault(
                key, {"name": pupil, "groups": [], "grades": defaultdict(int)}
            )
            if name not in item["groups"]:
                item["groups"].append(name)
            if grade is not None:
                item["grades"][grade] += 1

    if not groups:
        raise ValueError("В файле нет групп")
    student_records = []
    for item in students.values():
        grades = item["grades"]
        grade = None
        if grades:
            best = max(grades.values())
            candidates = [value for value, count in grades.items() if count == best]
            if len(candidates) == 1:
                grade = candidates[0]
        student_records.append(
            StudentRecord(full_name=item["name"], grade=grade, groups=tuple(item["groups"]))
        )
    return RosterData(
        teachers=tuple(
            TeacherRecord(item["name"], tuple(item["subjects"])) for item in teachers.values()
        ),
        groups=tuple(groups),
        students=tuple(student_records),
    )


def combine_group_rosters(rosters: list[RosterData]) -> RosterData:
    """Combine exports from several schools into one complete import snapshot."""
    teachers: dict[str, dict] = {}
    groups: dict[str, GroupRecord] = {}
    students: dict[str, dict] = {}
    for roster in rosters:
        for teacher in roster.teachers:
            key = normalize_name(teacher.name)
            item = teachers.setdefault(key, {"name": teacher.name, "subjects": []})
            for subject in teacher.subjects:
                if subject not in item["subjects"]:
                    item["subjects"].append(subject)
        for group in roster.groups:
            previous = groups.get(group.name)
            if previous is not None and previous != group:
                raise ValueError(
                    f"Группа «{group.name}» встречается в нескольких файлах с разными данными"
                )
            groups[group.name] = group
        for student in roster.students:
            key = normalize_name(student.full_name)
            item = students.setdefault(
                key, {"name": student.full_name, "groups": [], "grades": Counter()}
            )
            for group in student.groups:
                if group not in item["groups"]:
                    item["groups"].append(group)
            if student.grade is not None:
                item["grades"][student.grade] += 1

    student_records = []
    for item in students.values():
        grades = item["grades"]
        grade = None
        if grades:
            best = max(grades.values())
            candidates = [value for value, count in grades.items() if count == best]
            if len(candidates) == 1:
                grade = candidates[0]
        student_records.append(
            StudentRecord(full_name=item["name"], grade=grade, groups=tuple(item["groups"]))
        )
    return RosterData(
        teachers=tuple(
            TeacherRecord(item["name"], tuple(item["subjects"])) for item in teachers.values()
        ),
        groups=tuple(groups.values()),
        students=tuple(student_records),
    )


def _scores(answer: str | None) -> tuple[float | None, ...]:
    result = []
    for part in str(answer or "").split(","):
        try:
            result.append(float(part.strip()))
        except ValueError:
            result.append(None)
    return tuple(result)


def _when(value) -> datetime | None:
    try:
        return datetime.fromisoformat(str(value)) if value else None
    except ValueError:
        return None


def read_legacy(data: bytes) -> LegacyData:
    """Read results from a database of the previous version (SQLite)."""
    if not data.startswith(SQLITE_MAGIC):
        raise ValueError("Это не файл базы прошлой версии (.db)")
    handle, path = tempfile.mkstemp(suffix=".db")
    try:
        with os.fdopen(handle, "wb") as file:
            file.write(data)
        connection = sqlite3.connect(f"{Path(path).as_uri()}?mode=ro", uri=True)
        try:
            tables = {
                row[0]
                for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")
            }
            if not {"student", "exam"} <= tables:
                raise ValueError("В файле нет таблиц учеников и работ: это не база прошлой версии")
            works = tuple(
                LegacyWork(
                    fio=" ".join(str(fio or "").split()),
                    subject_code=str(subject or "").strip(),
                    scores=_scores(answer),
                    comment=str(comment or "").strip(),
                )
                for fio, subject, answer, comment in connection.execute(
                    "SELECT s.fio, e.subject, e.answer, e.comment "
                    "FROM exam e JOIN student s ON s.id = e.id_student ORDER BY e.id"
                )
            )
            title = created_at = None
            if "probnik" in tables:
                row = connection.execute(
                    "SELECT name, created_at FROM probnik ORDER BY id DESC LIMIT 1"
                ).fetchone()
                if row:
                    title, created_at = (str(row[0]).strip() or None), _when(row[1])
        except sqlite3.DatabaseError as error:
            raise ValueError("Не удалось прочитать базу прошлой версии") from error
        finally:
            connection.close()
    finally:
        os.unlink(path)
    return LegacyData(title=title, created_at=created_at, works=works)
