import re
from dataclasses import dataclass
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.importers.record import StudentRecord, normalize_name


@dataclass(frozen=True)
class ExcelFormat:
    name: str = "ФИО"
    status: str = "Статус обучения"
    groups: str = "Активные группы"
    grade: str = "Класс обучения"
    external_id: str = "ID"
    active_statuses: tuple[str, ...] = ("активен", "обучается")
    inactive_statuses: tuple[str, ...] = ("архив", "не обучается", "неактивен")
    group_separator: str | None = None


_WEEKDAY = r"(?:пн|вт|ср|чт|пт|сб|вс)\.?"
_SPACE = r"[^\S\r\n]*"
_TIME = rf"\d{{1,2}}:\d{{2}}(?:{_SPACE}[-–—]{_SPACE}\d{{1,2}}:\d{{2}})?"
_DAY_SEPARATOR = re.compile(
    rf"(?<!\w)({_WEEKDAY}(?:{_SPACE}{_TIME})?){_SPACE}(?:,|\bи\b){_SPACE}"
    rf"(?={_WEEKDAY}(?=\s|\d|[,;–—-]|$))",
    re.IGNORECASE,
)


def parse_groups(value: str, separator: str | None = None) -> tuple[str, ...]:
    """Normalize weekday separators, then split the student's list of groups.

    Commas and 'и' between weekdays (with optional times) become hyphens.
    Remaining commas, semicolons and line breaks separate groups.
    An explicit separator overrides these rules.
    """
    if separator is not None:
        parts = re.split(separator, value)
    else:
        value = _DAY_SEPARATOR.sub(r"\1-", value)
        parts = re.split(r"[,;\r\n]+", value)
    return tuple(dict.fromkeys(" ".join(part.split()) for part in parts if part.strip()))


def read_excel(source: BinaryIO, layout: ExcelFormat = ExcelFormat()) -> list[StudentRecord]:
    """CRM adapter only: no database imports or writes."""
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise ValueError("Не удалось открыть XLSX") from error
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [normalize_name(str(v or "")) for v in next(rows, ())]
        required = (layout.name, layout.status, layout.groups, layout.grade)
        for header in (*required, layout.external_id):
            count = headers.count(normalize_name(header))
            if count > 1 or (header in required and count != 1):
                raise ValueError(f"Ожидается одна колонка «{header}»")
        result = []
        for number, values in enumerate(rows, start=2):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            row = dict(zip(headers, values))

            def cell(header):
                return row.get(normalize_name(header))

            try:
                if any(isinstance(v, str) and v.startswith("=") for v in values):
                    raise ValueError("Формулы не поддерживаются; экспортируйте значения")
                status = normalize_name(str(cell(layout.status) or ""))
                if status not in (*layout.active_statuses, *layout.inactive_statuses):
                    raise ValueError(f"Неизвестный статус обучения: {status!r}")
                external_id = cell(layout.external_id)
                result.append(
                    StudentRecord(
                        full_name=str(cell(layout.name) or ""),
                        external_id=str(external_id) if external_id is not None else None,
                        grade=cell(layout.grade) if cell(layout.grade) != "" else None,
                        is_active=status in layout.active_statuses,
                        groups=parse_groups(str(cell(layout.groups) or ""), layout.group_separator),
                    )
                )
            except ValueError as error:
                raise ValueError(f"Строка {number}: {error}") from error
        return result
    finally:
        workbook.close()
